from docx import Document
from docx.shared import Cm
from PIL import Image
from io import BytesIO
import json
import sys
import os
import base64

sys.stdout.reconfigure(encoding="utf-8")


def add_photo_to_document(doc, photo_data, cell):
    try:
        image_binary = base64.b64decode(photo_data.split(",")[1])
    except:
        return

    image_stream = BytesIO(image_binary)
    image = Image.open(image_stream)

    width_cm, height_cm, dpi = 18, 13.5, 96
    width_px, height_px = int((width_cm / 2.54) * dpi), int((height_cm / 2.54) * dpi)

    image.thumbnail((width_px, height_px))

    temp_image_path = f"temp_{os.getpid()}.jpg"
    image.save(temp_image_path)

    paragraph = cell.add_paragraph()
    paragraph.alignment = 1
    run = paragraph.add_run()
    run.add_picture(temp_image_path, width=Cm(width_cm), height=Cm(height_cm))

    os.remove(temp_image_path)


def generate_document(json_file):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    uploads_dir = os.path.join(script_dir, "..", "uploads", "reports")
    os.makedirs(uploads_dir, exist_ok=True)
    template_path = os.path.join(script_dir, "template.docx")

    with open(json_file, "r", encoding="utf-8") as f:
        user_info = json.load(f)

    doc = Document(template_path)

    if isinstance(user_info.get("works", ""), list):
        user_info["works"] = "\n• " + "\n• ".join(user_info["works"])

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "[дата]" in cell.text:
                    cell.text = cell.text.replace("[дата]", user_info["date"])
                if "[адрес]" in cell.text:
                    cell.text = cell.text.replace("[адрес]", user_info["address"])
                if "[классификация]" in cell.text:
                    cell.text = cell.text.replace(
                        "[классификация]", user_info["classification"]
                    )
                if "[материалы]" in cell.text:
                    cell.text = cell.text.replace(
                        "[материалы]", user_info.get("material", "")
                    )
                if "[рекомендации]" in cell.text:
                    cell.text = cell.text.replace(
                        "[рекомендации]", user_info.get("recommendations", "")
                    )
                if "[дефекты]" in cell.text:
                    cell.text = cell.text.replace(
                        "[дефекты]", user_info.get("defects", "")
                    )
                if "[доп_работы]" in cell.text:
                    cell.text = cell.text.replace(
                        "[доп_работы]", user_info.get("additionalWorks", "")
                    )
                if "[комментарии]" in cell.text:
                    cell.text = cell.text.replace(
                        "[комментарии]", user_info.get("comments", "")
                    )
                if "[работы]" in cell.text:
                    checklist_text = "\n".join(
                        f"• {item['task']}"
                        for item in user_info.get("checklistItems", [])
                        if item.get("done")
                    )
                    cell.text = cell.text.replace(
                        "[работы]", checklist_text if checklist_text else ""
                    )
                if "[фио]" in cell.text:
                    full_name = f"{user_info['lastName']} {user_info['firstName']}"
                    cell.text = cell.text.replace("[фио]", full_name)
                if "[вставка]" in cell.text:
                    cell.text = ""
                    for photo in user_info.get("photos", []):
                        add_photo_to_document(doc, photo, cell)

    base_filename = f"Акт выполненных работ {user_info['date'].replace(':', '.')} {user_info['address']}.docx"
    output_path = os.path.join(uploads_dir, base_filename)

    counter = 1
    while os.path.exists(output_path):
        output_path = os.path.join(
            uploads_dir,
            f"Акт выполненных работ {user_info['date'].replace(':', '.')} {user_info['address']} ({counter}).docx",
        )
        counter += 1

    doc.save(output_path)

    print(output_path, "|", os.path.basename(output_path))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Ошибка: Не передан путь к JSON-файлу", file=sys.stderr)
        sys.exit(1)

    json_file_path = sys.argv[1]
    if not os.path.exists(json_file_path):
        print("Ошибка: JSON-файл не найден", file=sys.stderr)
        sys.exit(1)

    generate_document(json_file_path)
